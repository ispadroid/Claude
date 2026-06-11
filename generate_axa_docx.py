#!/usr/bin/env python3
"""Generate UPN_Wechsel_AXA.docx — AXA corporate one-pager."""

from docx import Document
from docx.shared import Cm, Pt, RGBColor, Twips
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import copy

OUT = "/home/user/Claude/UPN_Wechsel_AXA.docx"
IMG = "/root/.claude/uploads/cb3da536-4c85-53a7-8105-b6b8a4960349/bf765434-IMG_4052.png"

# ── colours ────────────────────────────────────────────────────────────────
BLUE   = "00008F"
RED    = "FF1721"
DARK   = "343C3D"
SECBG  = "EEEEF8"
WARNBG = "FFF0F0"
WARNTX = "8F0000"
GREY   = "F5F5F5"
WHITE  = "FFFFFF"
BRDGR  = "CCCCCC"
BRDBL  = "BBBBE8"
HINT   = "8C9BA5"
SUBTTL = "9999CC"

# ── helpers ────────────────────────────────────────────────────────────────

def bg(cell, hex6):
    """Set cell background shading."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex6)
    # Remove existing shd if present
    for existing in tcPr.findall(qn("w:shd")):
        tcPr.remove(existing)
    tcPr.append(shd)


def _get_or_add_tblPr(table):
    """Get or create tblPr element on a table."""
    tbl = table._tbl
    tblPr = tbl.find(qn("w:tblPr"))
    if tblPr is None:
        tblPr = OxmlElement("w:tblPr")
        tbl.insert(0, tblPr)
    return tblPr


def no_brd(table):
    """Remove all table borders."""
    tblPr = _get_or_add_tblPr(table)
    # Remove existing tblBorders
    for existing in tblPr.findall(qn("w:tblBorders")):
        tblPr.remove(existing)
    tblBorders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "none")
        el.set(qn("w:sz"), "0")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "auto")
        tblBorders.append(el)
    tblPr.append(tblBorders)


def box_brd(table, color_hex, sz=4):
    """Outer border only, no inside borders."""
    tblPr = _get_or_add_tblPr(table)
    for existing in tblPr.findall(qn("w:tblBorders")):
        tblPr.remove(existing)
    tblBorders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color_hex)
        tblBorders.append(el)
    for side in ("insideH", "insideV"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "none")
        el.set(qn("w:sz"), "0")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "auto")
        tblBorders.append(el)
    tblPr.append(tblBorders)


def cmar(cell, top, bot, left, right):
    """Set cell margins in twips."""
    tcPr = cell._tc.get_or_add_tcPr()
    for existing in tcPr.findall(qn("w:tcMar")):
        tcPr.remove(existing)
    tcMar = OxmlElement("w:tcMar")
    for name, val in (("top", top), ("bottom", bot), ("left", left), ("right", right)):
        el = OxmlElement(f"w:{name}")
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        tcMar.append(el)
    tcPr.append(tcMar)


def valign(cell, v):
    """Set vertical alignment."""
    tcPr = cell._tc.get_or_add_tcPr()
    for existing in tcPr.findall(qn("w:vAlign")):
        tcPr.remove(existing)
    vAlign = OxmlElement("w:vAlign")
    vAlign.set(qn("w:val"), v)
    tcPr.append(vAlign)


def fmt(p, align=WD_ALIGN_PARAGRAPH.LEFT, sb=0, sa=0, li=None):
    """Format paragraph."""
    p.alignment = align
    p.paragraph_format.space_before = Pt(sb)
    p.paragraph_format.space_after = Pt(sa)
    if li is not None:
        p.paragraph_format.left_indent = li


def run(p, text, bold=False, italic=False, sz=9, col=None):
    """Add a run with Arial font and optional hex color."""
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    r.font.name = "Arial"
    r.font.size = Pt(sz)
    if col:
        h = col
        r.font.color.rgb = RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))
    return r


def gap(doc, pt=4):
    """Add a minimal-height spacer paragraph."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    pPr = p._p.get_or_add_pPr()
    # Set line spacing to exact pt
    pSpacing = OxmlElement("w:spacing")
    pSpacing.set(qn("w:line"), str(int(pt * 20)))
    pSpacing.set(qn("w:lineRule"), "exact")
    # Remove existing spacing
    for existing in pPr.findall(qn("w:spacing")):
        pPr.remove(existing)
    pPr.append(pSpacing)
    # Make font tiny
    rPr = OxmlElement("w:rPr")
    sz_el = OxmlElement("w:sz")
    sz_el.set(qn("w:val"), "2")
    rPr.append(sz_el)
    pPr.append(rPr)
    return p


def cm_to_twips(cm_val):
    """Convert Cm() value (EMU) to twips. 1 twip = 635 EMU."""
    return int(int(cm_val) / 635)


def set_col_width(cell, width):
    """Set column width for a cell."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    for existing in tcPr.findall(qn("w:tcW")):
        tcPr.remove(existing)
    tcW = OxmlElement("w:tcW")
    tcW.set(qn("w:w"), str(cm_to_twips(width)))
    tcW.set(qn("w:type"), "dxa")
    tcPr.append(tcW)


def add_paragraph_shading(para, hex6):
    """Add background shading to a paragraph."""
    pPr = para._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex6)
    for existing in pPr.findall(qn("w:shd")):
        pPr.remove(existing)
    pPr.append(shd)


def add_paragraph_left_border(para, color_hex, sz=12, space=4):
    """Add a left border to a paragraph."""
    pPr = para._p.get_or_add_pPr()
    for existing in pPr.findall(qn("w:pBdr")):
        pPr.remove(existing)
    pBdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), str(sz))
    left.set(qn("w:space"), str(space))
    left.set(qn("w:color"), color_hex)
    pBdr.append(left)
    pPr.append(pBdr)


def add_top_border(para, color_hex, sz=4):
    """Add a top border to a paragraph."""
    pPr = para._p.get_or_add_pPr()
    for existing in pPr.findall(qn("w:pBdr")):
        pPr.remove(existing)
    pBdr = OxmlElement("w:pBdr")
    top = OxmlElement("w:top")
    top.set(qn("w:val"), "single")
    top.set(qn("w:sz"), str(sz))
    top.set(qn("w:space"), "4")
    top.set(qn("w:color"), color_hex)
    pBdr.append(top)
    pPr.append(pBdr)


def set_tbl_width(table, cm_val):
    """Set table width in twips."""
    tblPr = _get_or_add_tblPr(table)
    tblW = OxmlElement("w:tblW")
    tblW.set(qn("w:w"), str(cm_to_twips(cm_val)))
    tblW.set(qn("w:type"), "dxa")
    for existing in tblPr.findall(qn("w:tblW")):
        tblPr.remove(existing)
    tblPr.append(tblW)


# ── document setup ─────────────────────────────────────────────────────────

doc = Document()

# Page setup A4
section = doc.sections[0]
section.page_height = Cm(29.7)
section.page_width = Cm(21.0)
section.top_margin = Cm(1.0)
section.bottom_margin = Cm(1.0)
section.left_margin = Cm(1.2)
section.right_margin = Cm(1.2)

# Default styles
style = doc.styles["Normal"]
style.font.name = "Arial"
style.font.size = Pt(9)
pf = style.paragraph_format
pf.space_before = Pt(0)
pf.space_after = Pt(0)

# ── 1. HEADER TABLE ────────────────────────────────────────────────────────

hdr_tbl = doc.add_table(rows=1, cols=2)
hdr_tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
no_brd(hdr_tbl)
set_tbl_width(hdr_tbl, Cm(18.6))

hdr_c1 = hdr_tbl.cell(0, 0)
hdr_c2 = hdr_tbl.cell(0, 1)

# Set column widths
set_col_width(hdr_c1, Cm(15))
set_col_width(hdr_c2, Cm(3.6))

bg(hdr_c1, BLUE)
bg(hdr_c2, BLUE)
cmar(hdr_c1, 80, 80, 170, 60)
cmar(hdr_c2, 80, 80, 60, 120)
valign(hdr_c1, "center")
valign(hdr_c2, "center")

# Col1 Para1: AXA | UPN-Wechsel
p1 = hdr_c1.paragraphs[0]
fmt(p1)
run(p1, "AXA", bold=True, sz=18, col=WHITE)
run(p1, "  |  ", bold=False, sz=14, col=WHITE)
run(p1, "UPN-Wechsel — Kurzanleitung", bold=True, sz=14, col=WHITE)

# Col1 Para2: subtitle
p2 = hdr_c1.add_paragraph()
fmt(p2, sb=2)
run(p2, "Was nach der Änderung deiner geschäftlichen E-Mail-Adresse zu tun ist", sz=8.5, col=SUBTTL)

# Col2: AXA logo badge
p_logo = hdr_c2.paragraphs[0]
fmt(p_logo, align=WD_ALIGN_PARAGRAPH.RIGHT)
run(p_logo, "AXA", bold=True, sz=20, col=WHITE)

# ── 2. RED ACCENT BAR ──────────────────────────────────────────────────────

red_tbl = doc.add_table(rows=1, cols=1)
no_brd(red_tbl)
red_cell = red_tbl.cell(0, 0)
bg(red_cell, RED)
cmar(red_cell, 0, 0, 0, 0)
set_tbl_width(red_tbl, Cm(18.6))

# Set row height to ~75 twips
tr = red_tbl.rows[0]._tr
trPr = tr.find(qn("w:trPr"))
if trPr is None:
    trPr = OxmlElement("w:trPr")
    tr.insert(0, trPr)
trHeight = OxmlElement("w:trHeight")
trHeight.set(qn("w:val"), "75")
trHeight.set(qn("w:hRule"), "exact")
trPr.append(trHeight)

# Empty paragraph in red cell
p_red = red_cell.paragraphs[0]
fmt(p_red)
r_red = p_red.add_run(" ")
r_red.font.size = Pt(1)

# ── 3. GAP ─────────────────────────────────────────────────────────────────

gap(doc, 3)

# ── 4. INTRO TABLE ─────────────────────────────────────────────────────────

intro_tbl = doc.add_table(rows=1, cols=1)
box_brd(intro_tbl, BRDGR, sz=4)
set_tbl_width(intro_tbl, Cm(18.6))

intro_cell = intro_tbl.cell(0, 0)
bg(intro_cell, GREY)
cmar(intro_cell, 80, 80, 120, 120)

p_intro = intro_cell.paragraphs[0]
fmt(p_intro, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
run(p_intro, "Dein ", sz=8.5, col=DARK)
run(p_intro, "User Principal Name (UPN)", bold=True, sz=8.5, col=DARK)
run(p_intro, " — deine geschäftliche E-Mail-Adresse — wurde geändert. Einige Microsoft-365-Apps müssen manuell neu verknüpft werden. Folge den Schritten — es dauert nur wenige Minuten.", sz=8.5, col=DARK)

# ── 5. GAP ─────────────────────────────────────────────────────────────────

gap(doc, 4)

# ── 6. SECTION 1 HEADER ────────────────────────────────────────────────────

def section_header(doc, text):
    tbl = doc.add_table(rows=1, cols=1)
    no_brd(tbl)
    set_tbl_width(tbl, Cm(18.6))
    cell = tbl.cell(0, 0)
    bg(cell, BLUE)
    cmar(cell, 70, 70, 120, 120)
    p = cell.paragraphs[0]
    fmt(p)
    run(p, text, bold=True, sz=10.5, col=WHITE)
    return tbl


section_header(doc, "1 — Microsoft OneNote — Notizbücher neu verknüpfen")

# ── 7. SECTION 1 BODY ──────────────────────────────────────────────────────

s1_tbl = doc.add_table(rows=1, cols=2)
box_brd(s1_tbl, BRDBL, sz=4)
set_tbl_width(s1_tbl, Cm(18.6))

s1_left = s1_tbl.cell(0, 0)
s1_right = s1_tbl.cell(0, 1)
set_col_width(s1_left, Cm(11.0))
set_col_width(s1_right, Cm(7.6))

bg(s1_left, SECBG)
bg(s1_right, SECBG)
cmar(s1_left, 80, 80, 120, 80)
cmar(s1_right, 80, 80, 80, 80)
valign(s1_left, "top")
valign(s1_right, "center")

# Steps for section 1
steps_s1 = [
    ("⚠", "Vor dem Schließen: Sync-Fehler prüfen",
     "Sicherstellen, dass keine Sync-Fehler vorhanden sind. Falls ja: betroffene Notizbücher zuerst manuell aus dem lokalen Ordner sichern.",
     RED, True),
    ("1", "Alle Notizbücher schließen",
     "Rechtsklick auf jedes Notizbuch im linken Bereich → Notizbuch schließen (Close This Notebook). Für alle wiederholen, dann App beenden.",
     BLUE, False),
    ("2", "Bei OneDrive Web mit neuer Adresse anmelden",
     "Browser öffnen → onedrive.com → mit der neuen geschäftlichen E-Mail-Adresse (neuem UPN) anmelden.",
     BLUE, False),
    ("3", "Notizbuch direkt aus OneDrive Web öffnen",
     "Zum Ordner des Notizbuchs navigieren und darauf klicken: öffnet sich in OneNote für das Web. (Erzwingt die Neuverknüpfung mit dem neuen UPN.)",
     BLUE, False),
    ("4", "In der Desktop-App weiterarbeiten",
     "In OneNote für das Web oben rechts auf «In OneNote öffnen» klicken. Das Notizbuch ist nun korrekt mit dem neuen Konto verknüpft.",
     BLUE, False),
]

first_para = True
for num, label, desc, num_col, is_warning in steps_s1:
    # Label para
    if first_para:
        p_label = s1_left.paragraphs[0]
        first_para = False
    else:
        p_label = s1_left.add_paragraph()
    fmt(p_label, sb=4)
    run(p_label, num + "  ", bold=True, sz=12, col=num_col)
    run(p_label, label, bold=True, sz=9, col=RED if is_warning else DARK)

    # Description para
    p_desc = s1_left.add_paragraph()
    fmt(p_desc, li=Cm(0.6), sa=2)
    desc_col = RED if is_warning else DARK
    run(p_desc, desc, sz=8, col=desc_col)

# Right col: image + caption
p_img = s1_right.paragraphs[0]
fmt(p_img, align=WD_ALIGN_PARAGRAPH.CENTER)

img_width = Cm(7.6) - Cm(1.3)
img_height = img_width * (351 / 345)

try:
    r_img = p_img.add_run()
    r_img.add_picture(IMG, width=img_width, height=img_height)
except Exception as e:
    run(p_img, f"[Bild: {e}]", sz=7.5, col=HINT)

p_cap = s1_right.add_paragraph()
fmt(p_cap, align=WD_ALIGN_PARAGRAPH.CENTER, sb=4)
run(p_cap, "Rechtsklick → Notizbuch schließen", italic=True, sz=7.5, col=HINT)

# ── 8. GAP ─────────────────────────────────────────────────────────────────

gap(doc, 4)

# ── 9. SECTION 2 HEADER ────────────────────────────────────────────────────

section_header(doc, "2 — OneDrive — Freigaben neu erstellen")

# ── 10. SECTION 2 BODY ─────────────────────────────────────────────────────

s2_tbl = doc.add_table(rows=1, cols=1)
box_brd(s2_tbl, BRDBL, sz=4)
set_tbl_width(s2_tbl, Cm(18.6))

s2_cell = s2_tbl.cell(0, 0)
bg(s2_cell, SECBG)
cmar(s2_cell, 80, 80, 120, 120)

# Warning para
p_warn = s2_cell.paragraphs[0]
fmt(p_warn, li=Cm(0.1), sa=5)
add_paragraph_shading(p_warn, WARNBG)
add_paragraph_left_border(p_warn, RED, sz=12, space=4)
run(p_warn, "⚠  ", bold=True, sz=9, col=RED)
run(p_warn, "Alle Freigaben mit dem alten UPN werden ", sz=8.5, col=WARNTX)
run(p_warn, "automatisch ungültig.", bold=True, sz=8.5, col=WARNTX)
run(p_warn, " Empfänger verlieren den Zugriff — jedes Element muss manuell neu geteilt werden.", sz=8.5, col=WARNTX)

# Steps for section 2
steps_s2 = [
    ("1", "Freigegebene Elemente identifizieren",
     "In OneDrive Web oben auf «Geteilt» klicken → «Von mir geteilt» auswählen. Alle Elemente notieren, die mit dem alten UPN geteilt wurden."),
    ("2", "Freigaben entfernen und neu erstellen",
     "Für jedes Element: Rechtsklick → «Freigabe verwalten» → bestehende Freigaben entfernen → neu teilen mit der neuen E-Mail-Adresse der Empfänger."),
    ("3", "Empfänger informieren",
     "Empfänger über die neuen Freigabe-Links informieren. Die alten Links funktionieren nicht mehr."),
]

for num, label, desc in steps_s2:
    p_label = s2_cell.add_paragraph()
    fmt(p_label, sb=4)
    run(p_label, num + "  ", bold=True, sz=12, col=BLUE)
    run(p_label, label, bold=True, sz=9, col=DARK)

    p_desc = s2_cell.add_paragraph()
    fmt(p_desc, li=Cm(0.6), sa=2)
    run(p_desc, desc, sz=8, col=DARK)

# ── 11. GAP ────────────────────────────────────────────────────────────────

gap(doc, 4)

# ── 12. FOOTER ─────────────────────────────────────────────────────────────

p_footer = doc.add_paragraph()
fmt(p_footer, align=WD_ALIGN_PARAGRAPH.CENTER, sb=4)
add_top_border(p_footer, BRDGR, sz=4)
run(p_footer, "Bei Fragen wende dich an das IT-Team  ·  Internes Dokument", sz=7.5, col=HINT)

# ── SAVE ───────────────────────────────────────────────────────────────────

doc.save(OUT)
print(f"✓ Saved: {OUT}")
