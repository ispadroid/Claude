#!/usr/bin/env python3
"""
Generate AXA UPN-Wechsel Kurzanleitung PPTX (A4 portrait, single slide)
"""

from pptx import Presentation
from pptx.util import Cm, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from lxml import etree
import os

# ── Colours ─────────────────────────────────────────────────────────────────
AXA_BLUE   = RGBColor(0x00, 0x00, 0x8F)
AXA_RED    = RGBColor(0xFF, 0x17, 0x21)
TEXT_DARK  = RGBColor(0x34, 0x3C, 0x3D)
LIGHT_BLUE = RGBColor(0xEE, 0xEE, 0xF8)
LIGHT_RED  = RGBColor(0xFF, 0xF0, 0xF0)
WHITE      = RGBColor(0xFF, 0xFF, 0xFF)
GREY_HINT  = RGBColor(0x8C, 0x9B, 0xA5)
GREY_BG    = RGBColor(0xF5, 0xF5, 0xF5)
GREY_BORD  = RGBColor(0xCC, 0xCC, 0xCC)
BLUE_BORD  = RGBColor(0xBB, 0xBB, 0xE8)
DARK_RED   = RGBColor(0x8F, 0x00, 0x00)

# ── Page constants ───────────────────────────────────────────────────────────
PAGE_W = Cm(21)
PAGE_H = Cm(29.7)
ML     = Cm(1.1)
MR     = Cm(1.1)
CW     = PAGE_W - ML - MR   # Cm(18.8)

# ── Low-level helpers ────────────────────────────────────────────────────────

def add_rect(slide, x, y, w, h, fill_rgb=None, border_rgb=None, border_pt=0.5):
    shape = slide.shapes.add_shape(1, int(x), int(y), int(w), int(h))
    if fill_rgb:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill_rgb
    else:
        shape.fill.background()
    if border_rgb:
        shape.line.color.rgb = border_rgb
        shape.line.width = Pt(border_pt)
    else:
        shape.line.fill.background()
    return shape


def add_textbox(slide, x, y, w, h):
    tb = slide.shapes.add_textbox(int(x), int(y), int(w), int(h))
    tf = tb.text_frame
    tf.word_wrap = True
    return tb, tf


def set_margins(tb, l=Cm(0.1), t=Cm(0.04), r=Cm(0.1), b=Cm(0.04)):
    bp = tb.text_frame._txBody.find(qn('a:bodyPr'))
    if bp is not None:
        bp.set('lIns', str(int(l)))
        bp.set('rIns', str(int(r)))
        bp.set('tIns', str(int(t)))
        bp.set('bIns', str(int(b)))


def set_zero_margins(tb):
    set_margins(tb, Cm(0), Cm(0), Cm(0), Cm(0))


def para_spacing(para, before_pt=0, after_pt=0):
    pPr = para._p.get_or_add_pPr()
    for tag in (qn('a:spcBef'), qn('a:spcAft')):
        for el in pPr.findall(tag):
            pPr.remove(el)
    if before_pt:
        sb = etree.SubElement(pPr, qn('a:spcBef'))
        etree.SubElement(sb, qn('a:spcPts')).set('val', str(int(before_pt * 100)))
    if after_pt:
        sa = etree.SubElement(pPr, qn('a:spcAft'))
        etree.SubElement(sa, qn('a:spcPts')).set('val', str(int(after_pt * 100)))


def line_spacing_pct(para, pct=100):
    pPr = para._p.get_or_add_pPr()
    for el in pPr.findall(qn('a:lnSpc')):
        pPr.remove(el)
    ls = etree.SubElement(pPr, qn('a:lnSpc'))
    etree.SubElement(ls, qn('a:spcPct')).set('val', str(int(pct * 1000)))


def run(para, text, size_pt=9, bold=False, italic=False, color=TEXT_DARK):
    r = para.add_run()
    r.text = text
    r.font.name = 'Arial'
    r.font.size = Pt(size_pt)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    return r


# ── Composite helper: one step row ──────────────────────────────────────────

def draw_step(slide, x, y, w, h, num_text, num_color,
              label, body_text, body_color=TEXT_DARK, label_color=TEXT_DARK):
    NUM_W = Cm(0.72)
    GAP   = Cm(0.12)
    TXT_W = w - NUM_W - GAP

    # Number
    nb, ntf = add_textbox(slide, x, y + Cm(0.04), NUM_W, h)
    set_zero_margins(nb)
    ntf.word_wrap = False
    p = ntf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run(p, num_text, size_pt=12, bold=True, color=num_color)

    # Label + body text
    tb, tf = add_textbox(slide, x + NUM_W + GAP, y, TXT_W, h)
    set_zero_margins(tb)
    tf.word_wrap = True

    p1 = tf.paragraphs[0]
    p1.alignment = PP_ALIGN.LEFT
    para_spacing(p1, 0, 1)
    run(p1, label, size_pt=8.5, bold=True, color=label_color)

    p2 = tf.add_paragraph()
    p2.alignment = PP_ALIGN.LEFT
    para_spacing(p2, 0, 0)
    line_spacing_pct(p2, 100)
    run(p2, body_text, size_pt=7.8, bold=False, color=body_color)


# ════════════════════════════════════════════════════════════════════════════
# BUILD
# ════════════════════════════════════════════════════════════════════════════

prs = Presentation()
prs.slide_width  = PAGE_W
prs.slide_height = PAGE_H

slide = prs.slides.add_slide(prs.slide_layouts[6])   # blank

# Remove any leftover placeholders
for ph in list(slide.placeholders):
    ph._element.getparent().remove(ph._element)

# ── Vertical geometry ────────────────────────────────────────────────────────
TOP      = Cm(0.55)
H_HDR    = Cm(2.0)
H_RED    = Cm(0.09)
H_INTRO  = Cm(0.95)
H_SHDR   = Cm(0.58)    # section header height (used twice)
H_FOOT   = Cm(0.65)
VGAP     = Cm(0.12)

fixed = (TOP + H_HDR + H_RED + VGAP +
         H_INTRO + VGAP +
         H_SHDR + VGAP +          # sec1 header
         H_SHDR + VGAP +          # sec2 header
         VGAP + H_FOOT)

remaining   = PAGE_H - fixed
H_S1_BODY   = int(remaining * 0.55)
H_S2_BODY   = int(remaining - H_S1_BODY - VGAP)

y_hdr    = TOP
y_red    = y_hdr  + H_HDR
y_intro  = y_red  + H_RED  + VGAP
y_s1hdr  = y_intro + H_INTRO + VGAP
y_s1bod  = y_s1hdr + H_SHDR
y_s2hdr  = y_s1bod  + H_S1_BODY + VGAP
y_s2bod  = y_s2hdr + H_SHDR
y_foot   = y_s2bod  + H_S2_BODY + VGAP

# ── 1. HEADER ────────────────────────────────────────────────────────────────
add_rect(slide, ML, y_hdr, CW, H_HDR, fill_rgb=AXA_BLUE)

tb, tf = add_textbox(slide, ML + Cm(0.35), y_hdr + Cm(0.22), CW - Cm(0.5), Cm(0.82))
set_zero_margins(tb)
p = tf.paragraphs[0]
p.alignment = PP_ALIGN.LEFT
run(p, 'AXA',                            size_pt=22, bold=True,  color=WHITE)
run(p, '  |  ',                          size_pt=14, bold=False, color=WHITE)
run(p, 'UPN-Wechsel — Kurzanleitung',    size_pt=14, bold=True,  color=WHITE)

tb2, tf2 = add_textbox(slide, ML + Cm(0.35), y_hdr + Cm(1.22), CW - Cm(0.5), Cm(0.65))
set_zero_margins(tb2)
p2 = tf2.paragraphs[0]
p2.alignment = PP_ALIGN.LEFT
run(p2,
    'Was nach der Änderung deiner geschäftlichen E-Mail-Adresse zu tun ist',
    size_pt=8.5, bold=False, color=WHITE)

# ── 2. RED BAR ───────────────────────────────────────────────────────────────
add_rect(slide, ML, y_red, CW, H_RED, fill_rgb=AXA_RED)

# ── 3. INTRO BOX ─────────────────────────────────────────────────────────────
add_rect(slide, ML, y_intro, CW, H_INTRO,
         fill_rgb=GREY_BG, border_rgb=GREY_BORD, border_pt=0.5)
tb, tf = add_textbox(slide,
                     ML + Cm(0.2), y_intro + Cm(0.06),
                     CW - Cm(0.4), H_INTRO - Cm(0.1))
set_margins(tb, Cm(0.1), Cm(0.0), Cm(0.1), Cm(0.0))
tf.word_wrap = True
p = tf.paragraphs[0]
p.alignment = PP_ALIGN.JUSTIFY
line_spacing_pct(p, 100)
run(p,
    'Dein User Principal Name (UPN) — deine geschäftliche E-Mail-Adresse — wurde geändert. '
    'Einige Microsoft-365-Apps müssen manuell neu verknüpft werden. '
    'Folge den Schritten — es dauert nur wenige Minuten.',
    size_pt=8.5, color=TEXT_DARK)

# ── 4. SECTION 1 HEADER ──────────────────────────────────────────────────────
add_rect(slide, ML, y_s1hdr, CW, H_SHDR, fill_rgb=AXA_BLUE)
tb, tf = add_textbox(slide, ML + Cm(0.3), y_s1hdr + Cm(0.08), CW - Cm(0.4), H_SHDR - Cm(0.08))
set_zero_margins(tb)
p = tf.paragraphs[0]
p.alignment = PP_ALIGN.LEFT
run(p, '1 — Microsoft OneNote — Notizbücher neu verknüpfen',
    size_pt=10.5, bold=True, color=WHITE)

# ── 5. SECTION 1 BODY ────────────────────────────────────────────────────────
add_rect(slide, ML, y_s1bod, CW, H_S1_BODY,
         fill_rgb=LIGHT_BLUE, border_rgb=BLUE_BORD, border_pt=0.6)

LEFT_W  = CW * 0.60
RIGHT_W = CW - LEFT_W - Cm(0.2)
X_L     = ML + Cm(0.15)
X_R     = ML + LEFT_W + Cm(0.2)

ITOP = Cm(0.18)
steps1 = [
    ('⚠', AXA_RED,
     'Vor dem Schließen: Sync-Fehler prüfen',
     'Sicherstellen, dass keine Sync-Fehler vorhanden sind. '
     'Falls ja: betroffene Notizbücher zuerst manuell aus dem lokalen Ordner sichern.',
     AXA_RED, AXA_RED),
    ('1', AXA_BLUE,
     'Alle Notizbücher schließen',
     'Rechtsklick auf jedes Notizbuch im linken Bereich → Notizbuch schließen '
     '(Close This Notebook). Für alle wiederholen, dann App beenden.',
     TEXT_DARK, TEXT_DARK),
    ('2', AXA_BLUE,
     'Bei OneDrive Web mit neuer Adresse anmelden',
     'Browser öffnen → onedrive.com → mit der neuen geschäftlichen '
     'E-Mail-Adresse (neuem UPN) anmelden.',
     TEXT_DARK, TEXT_DARK),
    ('3', AXA_BLUE,
     'Notizbuch direkt aus OneDrive Web öffnen',
     'Zum Ordner des Notizbuchs navigieren und darauf klicken: öffnet sich in '
     'OneNote für das Web. (Erzwingt die Neuverknüpfung mit dem neuen UPN.)',
     TEXT_DARK, TEXT_DARK),
    ('4', AXA_BLUE,
     'In der Desktop-App weiterarbeiten',
     'In OneNote für das Web oben rechts auf «In OneNote öffnen» klicken. '
     'Das Notizbuch ist nun korrekt mit dem neuen Konto verknüpft.',
     TEXT_DARK, TEXT_DARK),
]

N1      = len(steps1)
avail_h = H_S1_BODY - ITOP - Cm(0.1)
sh      = avail_h // N1

for i, (num, ncol, label, body, bcol, lcol) in enumerate(steps1):
    sy = y_s1bod + ITOP + i * sh
    draw_step(slide, X_L, sy, LEFT_W - Cm(0.3), sh - Cm(0.06),
              num, ncol, label, body, body_color=bcol, label_color=lcol)
    # thin separator
    if i < N1 - 1:
        add_rect(slide, X_L, sy + sh - Cm(0.025),
                 LEFT_W - Cm(0.35), Cm(0.012), fill_rgb=BLUE_BORD)

# Right column: screenshot
IMG = '/root/.claude/uploads/cb3da536-4c85-53a7-8105-b6b8a4960349/bf765434-IMG_4052.png'
MAX_IW = RIGHT_W - Cm(0.35)
MAX_IH = H_S1_BODY - Cm(0.65)
RATIO   = 345 / 351
if MAX_IW / RATIO <= MAX_IH:
    iw = MAX_IW;  ih = MAX_IW / RATIO
else:
    ih = MAX_IH;  iw = MAX_IH * RATIO

ix = X_R + (RIGHT_W - iw) / 2
iy = y_s1bod + Cm(0.18)

# White frame
FP = Cm(0.1)
add_rect(slide, int(ix - FP), int(iy - FP),
         int(iw + FP*2), int(ih + FP*2),
         fill_rgb=WHITE, border_rgb=BLUE_BORD, border_pt=0.5)

if os.path.exists(IMG):
    slide.shapes.add_picture(IMG, int(ix), int(iy), int(iw), int(ih))

# Caption
cap_y = int(iy + ih + FP + Cm(0.06))
tb, tf = add_textbox(slide, int(X_R), cap_y, int(RIGHT_W), Cm(0.4))
set_zero_margins(tb)
p = tf.paragraphs[0]
p.alignment = PP_ALIGN.CENTER
run(p, 'Rechtsklick → Notizbuch schließen', size_pt=7.5, italic=True, color=GREY_HINT)

# ── 6. SECTION 2 HEADER ──────────────────────────────────────────────────────
add_rect(slide, ML, y_s2hdr, CW, H_SHDR, fill_rgb=AXA_BLUE)
tb, tf = add_textbox(slide, ML + Cm(0.3), y_s2hdr + Cm(0.08), CW - Cm(0.4), H_SHDR - Cm(0.08))
set_zero_margins(tb)
p = tf.paragraphs[0]
p.alignment = PP_ALIGN.LEFT
run(p, '2 — OneDrive — Freigaben neu erstellen',
    size_pt=10.5, bold=True, color=WHITE)

# ── 7. SECTION 2 BODY ────────────────────────────────────────────────────────
add_rect(slide, ML, y_s2bod, CW, H_S2_BODY,
         fill_rgb=LIGHT_BLUE, border_rgb=BLUE_BORD, border_pt=0.6)

WPAD = Cm(0.12)
WARN_H = Cm(0.95)
warn_y = y_s2bod + WPAD
warn_x = ML + Cm(0.15)
warn_w = CW - Cm(0.3)

# Warning background (light red)
add_rect(slide, warn_x, warn_y, warn_w, WARN_H, fill_rgb=LIGHT_RED)
# Left red stripe
add_rect(slide, warn_x, warn_y, Cm(0.18), WARN_H, fill_rgb=AXA_RED)

# Warning text
tb, tf = add_textbox(slide,
                     warn_x + Cm(0.28), warn_y + Cm(0.08),
                     warn_w - Cm(0.35), WARN_H - Cm(0.14))
set_margins(tb, Cm(0.05), Cm(0), Cm(0.05), Cm(0))
tf.word_wrap = True
p = tf.paragraphs[0]
p.alignment = PP_ALIGN.LEFT
line_spacing_pct(p, 108)
run(p, '⚠  ',                                           size_pt=8.5, bold=True,  color=AXA_RED)
run(p, 'Alle Freigaben mit dem alten UPN werden ',      size_pt=8.5, bold=False, color=DARK_RED)
run(p, 'automatisch ungültig.',                          size_pt=8.5, bold=True,  color=DARK_RED)
run(p, ' Empfänger verlieren den Zugriff — jedes Element muss manuell neu geteilt werden.',
    size_pt=8.5, bold=False, color=DARK_RED)

# Section 2 steps
steps2 = [
    ('1',
     'Freigegebene Elemente finden',
     'Bei onedrive.com mit neuer Adresse anmelden → im linken Bereich auf '
     '«Geteilt» klicken → alle früheren Freigaben werden angezeigt.'),
    ('2',
     'Jede Datei / jeden Ordner neu freigeben',
     'Rechtsklick → Freigeben → Empfänger eingeben → Berechtigung wählen '
     '(Anzeigen / Bearbeiten) → Senden.'),
    ('3',
     'Empfänger über neuen Link informieren',
     'Alte Links funktionieren nicht mehr. Neue Einladung versenden oder Link '
     'direkt mitteilen. Für SharePoint-Bibliotheken das IT-Team kontaktieren.'),
]

s2_area_y = warn_y + WARN_H + Cm(0.14)
s2_area_h = H_S2_BODY - WPAD - WARN_H - Cm(0.22)
s2h       = s2_area_h // len(steps2)

for i, (num, label, body) in enumerate(steps2):
    sy = s2_area_y + i * s2h
    draw_step(slide, ML + Cm(0.15), sy, CW - Cm(0.3), s2h - Cm(0.05),
              num, AXA_BLUE, label, body)
    if i < len(steps2) - 1:
        add_rect(slide, ML + Cm(0.15), sy + s2h - Cm(0.025),
                 CW - Cm(0.35), Cm(0.012), fill_rgb=BLUE_BORD)

# ── 8. FOOTER ────────────────────────────────────────────────────────────────
add_rect(slide, ML, y_foot, CW, Cm(0.025), fill_rgb=GREY_BORD)
tb, tf = add_textbox(slide, ML, y_foot + Cm(0.06), CW, H_FOOT - Cm(0.06))
set_zero_margins(tb)
p = tf.paragraphs[0]
p.alignment = PP_ALIGN.CENTER
run(p, 'Bei Fragen wende dich an das IT-Team  ·  Internes Dokument',
    size_pt=7.5, color=GREY_HINT)

# ── Save ─────────────────────────────────────────────────────────────────────
OUT = '/home/user/Claude/UPN_Wechsel_AXA.pptx'
prs.save(OUT)

w_cm = prs.slide_width.cm
h_cm = prs.slide_height.cm
print(f'Saved: {OUT}')
print(f'Slide dimensions: {w_cm:.1f} cm x {h_cm:.1f} cm  (A4 portrait)')
print(f'Slides in file: {len(prs.slides)}')
print(f'Section 1 body height: {H_S1_BODY / 914400 * 2.54:.2f} cm')
print(f'Section 2 body height: {H_S2_BODY / 914400 * 2.54:.2f} cm')
