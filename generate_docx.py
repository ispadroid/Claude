from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUTPUT     = "/home/user/Claude/UPN_Wechsel_Handout.docx"
SCREENSHOT = "/root/.claude/uploads/cb3da536-4c85-53a7-8105-b6b8a4960349/bf765434-IMG_4052.png"

# ── Colours ───────────────────────────────────────────────────────────────────
CD    = "0F3D6E"   # blue dark
CM    = "1B6EC2"   # blue mid
CL    = "EAF2FD"   # blue light
CO    = "C45C00"   # orange
COB   = "FFF4EC"   # orange bg
COT   = "5C2800"   # orange text dark
CG    = "2D2D2D"   # grey text
CGl   = "F4F6F9"   # grey light
CW_   = "FFFFFF"   # white
CB    = "C0D4EE"   # border blue
CH    = "555555"   # hint grey
CBHNT = "BDD6F5"  # blue hint (header subtitle)

doc = Document()

# ── Page setup ────────────────────────────────────────────────────────────────
sec = doc.sections[0]
sec.page_width    = Cm(21)
sec.page_height   = Cm(29.7)
sec.top_margin    = Cm(1.0)
sec.bottom_margin = Cm(1.0)
sec.left_margin   = Cm(1.2)
sec.right_margin  = Cm(1.2)

CW = Cm(18.6)   # usable content width

# Suppress default paragraph spacing on Normal style
ns = doc.styles['Normal']
ns.paragraph_format.space_before = Pt(0)
ns.paragraph_format.space_after  = Pt(0)
ns.font.name = 'Calibri'
ns.font.size = Pt(9)

# ── XML helpers ───────────────────────────────────────────────────────────────

def bg(cell, hex6):
    tcPr = cell._tc.get_or_add_tcPr()
    for old in tcPr.findall(qn('w:shd')):
        tcPr.remove(old)
    s = OxmlElement('w:shd')
    s.set(qn('w:val'), 'clear')
    s.set(qn('w:color'), 'auto')
    s.set(qn('w:fill'), hex6)
    tcPr.append(s)

def no_brd(table):
    tblPr = table._tbl.tblPr
    for old in tblPr.findall(qn('w:tblBorders')):
        tblPr.remove(old)
    tb = OxmlElement('w:tblBorders')
    for n in ['top','left','bottom','right','insideH','insideV']:
        b = OxmlElement(f'w:{n}')
        b.set(qn('w:val'), 'none')
        b.set(qn('w:sz'), '0')
        b.set(qn('w:space'), '0')
        b.set(qn('w:color'), 'auto')
        tb.append(b)
    tblPr.append(tb)

def box_brd(table, col, sz=4):
    tblPr = table._tbl.tblPr
    for old in tblPr.findall(qn('w:tblBorders')):
        tblPr.remove(old)
    tb = OxmlElement('w:tblBorders')
    for n in ['top','left','bottom','right']:
        b = OxmlElement(f'w:{n}')
        b.set(qn('w:val'), 'single')
        b.set(qn('w:sz'), str(sz))
        b.set(qn('w:space'), '0')
        b.set(qn('w:color'), col)
        tb.append(b)
    for n in ['insideH','insideV']:
        b = OxmlElement(f'w:{n}')
        b.set(qn('w:val'), 'none')
        b.set(qn('w:sz'), '0')
        b.set(qn('w:space'), '0')
        b.set(qn('w:color'), 'auto')
        tb.append(b)
    tblPr.append(tb)

def cmar(cell, top=60, bot=60, left=120, right=120):
    tcPr = cell._tc.get_or_add_tcPr()
    m = OxmlElement('w:tcMar')
    for name, val in [('top',top),('bottom',bot),('left',left),('right',right)]:
        e = OxmlElement(f'w:{name}')
        e.set(qn('w:w'), str(val))
        e.set(qn('w:type'), 'dxa')
        m.append(e)
    tcPr.append(m)

def valign(cell, v='top'):
    tcPr = cell._tc.get_or_add_tcPr()
    e = OxmlElement('w:vAlign')
    e.set(qn('w:val'), v)
    tcPr.append(e)

def p_shd(para, fill_hex):
    pPr = para._p.get_or_add_pPr()
    s = OxmlElement('w:shd')
    s.set(qn('w:val'), 'clear')
    s.set(qn('w:color'), 'auto')
    s.set(qn('w:fill'), fill_hex)
    pPr.append(s)

def p_left_bar(para, col_hex, sz=12):
    pPr = para._p.get_or_add_pPr()
    bd = OxmlElement('w:pBdr')
    lb = OxmlElement('w:left')
    lb.set(qn('w:val'), 'single')
    lb.set(qn('w:sz'), str(sz))
    lb.set(qn('w:space'), '4')
    lb.set(qn('w:color'), col_hex)
    bd.append(lb)
    pPr.append(bd)

def p_top_brd(para, col='CCCCCC', sz=4):
    pPr = para._p.get_or_add_pPr()
    bd = OxmlElement('w:pBdr')
    tp = OxmlElement('w:top')
    tp.set(qn('w:val'), 'single')
    tp.set(qn('w:sz'), str(sz))
    tp.set(qn('w:space'), '1')
    tp.set(qn('w:color'), col)
    bd.append(tp)
    pPr.append(bd)

def fmt(p, align=WD_ALIGN_PARAGRAPH.LEFT, sb=0, sa=0, li=None):
    p.alignment = align
    p.paragraph_format.space_before = Pt(sb)
    p.paragraph_format.space_after  = Pt(sa)
    if li is not None:
        p.paragraph_format.left_indent = li

def run(p, text, bold=False, italic=False, sz=9, col=None):
    r = p.add_run(text)
    r.bold = bold; r.italic = italic
    r.font.size = Pt(sz)
    r.font.name = 'Calibri'
    if col:
        r.font.color.rgb = RGBColor.from_string(col)
    return r

def gap(doc, pt=3):
    """Tiny spacer paragraph."""
    p = doc.add_paragraph()
    fmt(p)
    pPr = p._p.get_or_add_pPr()
    rPr = OxmlElement('w:rPr')
    for tag in ['w:sz', 'w:szCs']:
        e = OxmlElement(tag)
        e.set(qn('w:val'), str(int(pt * 2)))
        rPr.append(e)
    pPr.append(rPr)

# ════════════════════════════════════════════════════════════════════════════════
# HEADER
# ════════════════════════════════════════════════════════════════════════════════
hdr = doc.add_table(rows=1, cols=2)
no_brd(hdr)
hdr.columns[0].width = Cm(14.5)
hdr.columns[1].width = Cm(4.1)

c0, c1 = hdr.cell(0, 0), hdr.cell(0, 1)
for c in (c0, c1):
    bg(c, CD)
    valign(c, 'center')
cmar(c0, top=80, bot=80, left=170, right=60)
cmar(c1, top=80, bot=80, left=60, right=120)

p = c0.paragraphs[0]; fmt(p)
run(p, "UPN-Wechsel — Kurzanleitung", True, False, 17, CW_)
p2 = c0.add_paragraph(); fmt(p2, sb=2)
run(p2, "Was nach der Änderung deiner geschäftlichen E-Mail-Adresse zu tun ist",
    False, False, 8.5, CBHNT)

pb = c1.paragraphs[0]; fmt(pb, WD_ALIGN_PARAGRAPH.RIGHT)
run(pb, "Microsoft 365", True, False, 10.5, CBHNT)

# Orange accent bar
acc = doc.add_table(rows=1, cols=1)
no_brd(acc)
ac = acc.cell(0, 0)
bg(ac, CO)
cmar(ac, 0, 0, 0, 0)
tr = acc.rows[0]._tr
trPr = tr.get_or_add_trPr()
h = OxmlElement('w:trHeight')
h.set(qn('w:val'), '75')
h.set(qn('w:hRule'), 'exact')
trPr.append(h)
fmt(ac.paragraphs[0])

gap(doc, 3)

# ── Intro ─────────────────────────────────────────────────────────────────────
it = doc.add_table(rows=1, cols=1)
box_brd(it, CB, 4)
ic = it.cell(0, 0)
bg(ic, CGl)
cmar(ic, 60, 60, 150, 150)
pi = ic.paragraphs[0]; fmt(pi, WD_ALIGN_PARAGRAPH.JUSTIFY)
run(pi, "Dein ", sz=8.5, col=CG)
run(pi, "User Principal Name (UPN)", True, sz=8.5, col=CG)
run(pi, " — deine geschäftliche E-Mail-Adresse — wurde geändert. "
    "Einige Microsoft-365-Apps müssen manuell neu verknüpft werden. "
    "Folge den Schritten — es dauert nur wenige Minuten.", sz=8.5, col=CG)

gap(doc, 4)

# ════════════════════════════════════════════════════════════════════════════════
# SECTION 1 — ONENOTE
# ════════════════════════════════════════════════════════════════════════════════
def sec_hdr(title):
    t = doc.add_table(rows=1, cols=1)
    no_brd(t)
    c = t.cell(0, 0)
    bg(c, CM)
    cmar(c, 50, 50, 150, 120)
    p = c.paragraphs[0]; fmt(p)
    run(p, title, True, sz=10.5, col=CW_)

sec_hdr("1 — Microsoft OneNote — Notizbücher neu verknüpfen")

# Body: left steps | right screenshot
LW = Cm(11.0)
RW = CW - LW
s1b = doc.add_table(rows=1, cols=2)
box_brd(s1b, CB, 4)
s1b.columns[0].width = LW
s1b.columns[1].width = RW
sl = s1b.cell(0, 0)
sr = s1b.cell(0, 1)
for c in (sl, sr):
    bg(c, CL)
    valign(c, 'top')
cmar(sl, 80, 80, 130, 80)
cmar(sr, 80, 80, 80, 120)

steps1 = [
    ("⚠", "Vor dem Schließen: Sync-Fehler prüfen",
     "Sicherstellen, dass keine Sync-Fehler vorhanden sind. Falls ja: betroffene Notizbücher "
     "zuerst manuell aus dem lokalen Ordner sichern.",
     CO, CO),
    ("1", "Alle Notizbücher schließen",
     "Rechtsklick auf jedes Notizbuch im linken Bereich → Notizbuch schließen "
     "(Close This Notebook). Für alle wiederholen, dann App beenden.",
     CM, CG),
    ("2", "Bei OneDrive Web mit neuer Adresse anmelden",
     "Browser öffnen → onedrive.com → mit der neuen geschäftlichen "
     "E-Mail-Adresse (neuem UPN) anmelden.",
     CM, CG),
    ("3", "Notizbuch direkt aus OneDrive Web öffnen",
     "Zum Ordner des Notizbuchs navigieren und darauf klicken: öffnet sich "
     "in OneNote für das Web. (Erzwingt die Neuverknüpfung mit dem neuen UPN.)",
     CM, CG),
    ("4", "In der Desktop-App weiterarbeiten",
     "In OneNote für das Web oben rechts auf «In OneNote öffnen» klicken. "
     "Das Notizbuch ist nun korrekt mit dem neuen Konto verknüpft.",
     CM, CG),
]

first = True
for num, lbl, txt, nc, tc in steps1:
    ph = sl.paragraphs[0] if first else sl.add_paragraph()
    first = False
    fmt(ph, sb=4 if num != "⚠" else 0, sa=0)
    run(ph, f"{num}  ", True, sz=12, col=nc)
    run(ph, lbl, True, sz=9, col=CG)

    pt = sl.add_paragraph()
    fmt(pt, li=Cm(0.6), sa=2)
    run(pt, txt, sz=8, col=tc)

# Screenshot right
p_img = sr.paragraphs[0]
fmt(p_img, WD_ALIGN_PARAGRAPH.CENTER)
img_run = p_img.add_run()
IW = int(RW - Cm(1.3))
IH = int(IW * 351 / 345)
img_run.add_picture(SCREENSHOT, width=IW, height=IH)

p_cap = sr.add_paragraph()
fmt(p_cap, WD_ALIGN_PARAGRAPH.CENTER, sb=3)
run(p_cap, "Rechtsklick → Notizbuch schließen", italic=True, sz=7.5, col=CH)

gap(doc, 4)

# ════════════════════════════════════════════════════════════════════════════════
# SECTION 2 — ONEDRIVE
# ════════════════════════════════════════════════════════════════════════════════
sec_hdr("2 — OneDrive — Freigaben neu erstellen")

s2b = doc.add_table(rows=1, cols=1)
box_brd(s2b, CB, 4)
s2c = s2b.cell(0, 0)
bg(s2c, CL)
cmar(s2c, 80, 80, 130, 130)

# Warning paragraph
pw = s2c.paragraphs[0]
fmt(pw, sa=5, li=Cm(0.1))
p_shd(pw, COB)
p_left_bar(pw, CO, 12)
run(pw, "⚠  ", True, sz=9, col=CO)
run(pw, "Alle Freigaben mit dem alten UPN werden ", sz=8.5, col=COT)
run(pw, "automatisch ungültig.", True, sz=8.5, col=COT)
run(pw, " Empfänger verlieren den Zugriff — "
    "jedes Element muss manuell neu geteilt werden.", sz=8.5, col=COT)

steps2 = [
    ("1", "Freigegebene Elemente finden",
     "Bei onedrive.com mit neuer Adresse anmelden → im linken Bereich auf «Geteilt» klicken → "
     "alle früheren Freigaben werden angezeigt."),
    ("2", "Jede Datei / jeden Ordner neu freigeben",
     "Rechtsklick → Freigeben → Empfänger eingeben → Berechtigung wählen "
     "(Anzeigen / Bearbeiten) → Senden."),
    ("3", "Empfänger über neuen Link informieren",
     "Alte Links funktionieren nicht mehr. Neue Einladung versenden oder Link direkt mitteilen. "
     "Für SharePoint-Bibliotheken das IT-Team kontaktieren."),
]

for num, lbl, txt in steps2:
    ph = s2c.add_paragraph(); fmt(ph, sb=5, sa=0)
    run(ph, f"{num}  ", True, sz=12, col=CM)
    run(ph, lbl, True, sz=9, col=CG)
    pt = s2c.add_paragraph(); fmt(pt, li=Cm(0.6), sa=2)
    run(pt, txt, sz=8, col=CG)

# ── Footer ────────────────────────────────────────────────────────────────────
gap(doc, 4)
pf = doc.add_paragraph()
fmt(pf, WD_ALIGN_PARAGRAPH.CENTER, sb=2)
p_top_brd(pf)
run(pf, "Bei Fragen wende dich an das IT-Team  ·  Internes Dokument",
    sz=7.5, col=CH)

doc.save(OUTPUT)
print(f"DOCX erstellt: {OUTPUT}")
